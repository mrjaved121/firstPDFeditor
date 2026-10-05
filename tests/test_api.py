"""End-to-end tests of the server API. Run with:  python -m unittest discover tests"""

import base64
import io
import json
import sys
import tempfile
import unicodedata
import unittest
import zipfile
from pathlib import Path

import pymupdf
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import app as server  # noqa: E402


def sample_pdf():
    """Two pages: text with sensitive data and an image; the second page rotated."""
    doc = pymupdf.open()
    p = doc.new_page(width=595, height=842)
    p.insert_text((72, 100), "Invoice for Jane Doe", fontsize=16)
    p.insert_text((72, 130), "SSN: 123-45-6789", fontsize=12)
    p.insert_text((72, 150), "Card: 4111 1111 1111 1111", fontsize=12)
    p.insert_text((72, 170), "Email: jane.doe@example.com", fontsize=12)
    p.insert_text((72, 190), "Phone: (212) 555-0100", fontsize=12)
    p.insert_text((72, 210), "Not a card: 1234 5678 9012 3456", fontsize=12)
    p.insert_image(pymupdf.Rect(300, 300, 400, 350), stream=png_bytes())
    q = doc.new_page(width=595, height=842)
    q.insert_text((72, 100), "Second page", fontsize=14)
    q.set_rotation(90)
    doc.set_metadata({"author": "Secret Author", "producer": "Secret Producer", "title": "Doc"})
    return doc.tobytes()


def form_pdf():
    doc = pymupdf.open()
    p = doc.new_page()
    p.insert_text((72, 95), "Full name:", fontsize=11)
    p.insert_text((72, 135), "Date:", fontsize=11)
    for i, name in enumerate(("full_name", "date_signed")):
        w = pymupdf.Widget()
        w.field_type = pymupdf.PDF_WIDGET_TYPE_TEXT
        w.field_name = name
        w.rect = pymupdf.Rect(150, 80 + 40 * i, 350, 100 + 40 * i)
        p.add_widget(w)
    return doc.tobytes()


def png_bytes(size=(40, 20)):
    im = Image.new("RGB", size, "white")
    for x in range(10):
        for y in range(5):
            im.putpixel((x, y), (255, 0, 0))
    b = io.BytesIO()
    im.save(b, "PNG")
    return b.getvalue()


def text_of(page):
    """Page text with Arabic presentation forms turned back into ordinary letters."""
    return unicodedata.normalize("NFKC", page.get_text())


def data_url(data, mime="image/png"):
    return f"data:{mime};base64," + base64.b64encode(data).decode()


class ApiTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        server.WORKSPACE = Path(self.tmp.name)
        server.DOCS.clear()
        self.c = server.app.test_client()

    def tearDown(self):
        server.DOCS.clear()
        self.tmp.cleanup()

    # -------------------------------------------------------- helpers

    def open(self, data=None, name="test.pdf"):
        r = self.c.post("/api/open", data={"file": (io.BytesIO(data or sample_pdf()), name)},
                        content_type="multipart/form-data")
        self.assertEqual(r.status_code, 200, r.get_json())
        return r.get_json()

    def op(self, doc_id, ok=True, **body):
        r = self.c.post(f"/api/doc/{doc_id}/op", json=body)
        if ok:
            self.assertEqual(r.status_code, 200, r.get_json())
        return r

    def pdf(self, doc_id):
        return pymupdf.open(stream=server.current_bytes(server.DOCS[doc_id]), filetype="pdf")

    # -------------------------------------------------------- tests

    def test_open_tabs_and_autosave(self):
        a = self.open()
        b = self.open(form_pdf(), "form.pdf")
        ids = [d["id"] for d in self.c.get("/api/docs").get_json()]
        self.assertEqual(ids, [a["id"], b["id"]])
        self.assertTrue((server.WORKSPACE / f"{a['id']}.pdf").exists())
        # A restart restores both documents from the workspace.
        server.DOCS.clear()
        server.restore_workspace()
        self.assertEqual(set(server.DOCS), {a["id"], b["id"]})
        self.assertEqual(server.DOCS[b["id"]]["name"], "form.pdf")
        self.c.delete(f"/api/doc/{a['id']}")
        self.assertFalse((server.WORKSPACE / f"{a['id']}.pdf").exists())
        self.assertEqual(len(self.c.get("/api/docs").get_json()), 1)

    def test_add_text_styles(self):
        d = self.open()["id"]
        self.op(d, op="add_text", page=0, x=72, y=400, text="Plain centered\nline two", size=14,
                align="center", font="auto")
        r = self.op(d, op="add_text", page=0, x=72, y=450, text="Bold italic", size=14,
                    bold=True, italic=True, font="auto")
        self.assertIn("Bold Italic", r.get_json()["message"])
        r = self.op(d, op="add_text", page=0, x=72, y=500, text="مرحبا بالعالم", size=14, align="right")
        page = self.pdf(d)[0]
        self.assertTrue(any(a.type[1] == "FreeText" for a in page.annots()))
        self.assertIn("Bold italic", page.get_text())
        self.assertIn("مرحبا", text_of(page))

    def test_edit_text(self):
        d = self.open()["id"]
        lines = self.c.get(f"/api/doc/{d}/page/0/lines").get_json()
        line = next(l for l in lines if l["text"].startswith("Invoice"))
        self.op(d, op="edit_text", page=0, pageBbox=line["pageBbox"], text="Invoice for John Roe")
        text = self.pdf(d)[0].get_text()
        self.assertIn("John Roe", text)
        self.assertNotIn("Jane Doe", text)

    def test_markup_and_notes(self):
        d = self.open()["id"]
        for kind in ("highlight", "underline", "strikeout"):
            self.op(d, op=kind, page=0, rect=[70, 85, 300, 105])
        self.op(d, op="note", page=0, point=[500, 100], text="Check this")
        self.op(d, op="rect", page=0, rect=[50, 600, 150, 650], color="#ff0000", width=2)
        self.op(d, op="ellipse", page=0, rect=[200, 600, 300, 650])
        self.op(d, op="line", page=0, p1=[50, 700], p2=[200, 720], arrow=True)
        self.op(d, op="ink", page=0, paths=[[[50, 750], [60, 760], [80, 755]]])
        types = sorted(a.type[1] for a in self.pdf(d)[0].annots())
        self.assertEqual(types, sorted(["Highlight", "Underline", "StrikeOut", "Text", "Square",
                                        "Circle", "Line", "Ink"]))
        self.op(d, ok=False, op="note", page=0, point=[1, 1], text="  ")

    def test_redaction_removes_text(self):
        d = self.open()["id"]
        found = self.c.post(f"/api/doc/{d}/find", json={"kinds": ["ssn", "card", "email", "phone"]}).get_json()
        kinds = sorted(m["kind"] for m in found)
        self.assertEqual(kinds, ["card", "email", "phone", "ssn"])  # the non-Luhn number is ignored
        self.op(d, op="redact_areas", items=found)
        text = self.pdf(d)[0].get_text()
        for gone in ("123-45-6789", "4111", "jane.doe@example.com", "555-0100"):
            self.assertNotIn(gone, text)
        self.assertIn("Invoice", text)
        custom = self.c.post(f"/api/doc/{d}/find", json={"kinds": [], "custom": "invoice"}).get_json()
        self.assertEqual(len(custom), 1)
        bad = self.c.post(f"/api/doc/{d}/find", json={"kinds": [], "custom": "(", "regex": True})
        self.assertEqual(bad.status_code, 400)

    def test_whiteout_and_blackout(self):
        d = self.open()["id"]
        self.op(d, op="erase_area", page=0, rect=[60, 115, 300, 135], fill="#000000")
        self.assertNotIn("123-45-6789", self.pdf(d)[0].get_text())

    def test_images(self):
        d = self.open()["id"]
        imgs = self.c.get(f"/api/doc/{d}/page/0/images").get_json()
        self.assertEqual(len(imgs), 1)
        img = imgs[0]
        # Move, resize, turn and crop.
        self.op(d, op="image_transform", page=0, xref=img["xref"], bbox=img["bbox"],
                rect=[100, 500, 140, 580], rotate=90, crop=[0, 0, 0.5, 1])
        imgs = self.c.get(f"/api/doc/{d}/page/0/images").get_json()
        self.assertEqual(len(imgs), 1)
        self.assertEqual([round(v) for v in imgs[0]["bbox"]], [100, 500, 140, 580])
        self.assertEqual((imgs[0]["width"], imgs[0]["height"]), (20, 20))
        self.op(d, op="image_delete", page=0, xref=imgs[0]["xref"], bbox=imgs[0]["bbox"])
        self.assertEqual(self.c.get(f"/api/doc/{d}/page/0/images").get_json(), [])
        # Place PNG and SVG images, including on the rotated page.
        self.op(d, op="image", page=1, rect=[50, 50, 130, 90], dataUrl=data_url(png_bytes()))
        svg = b'<svg xmlns="http://www.w3.org/2000/svg" width="10" height="10"><rect width="10" height="10" fill="red"/></svg>'
        self.op(d, op="image", page=0, rect=[50, 50, 100, 100], dataUrl=data_url(svg, "image/svg+xml"))
        pix = self.pdf(d)[1].get_pixmap()
        self.assertEqual(pix.pixel(52, 52)[:3], (255, 0, 0))  # top-left red corner stays top-left
        self.op(d, ok=False, op="image", page=0, rect=[0, 0, 9, 9], dataUrl="data:image/svg+xml;base64,AAAA")

    def test_pages(self):
        d = self.open()["id"]
        self.op(d, op="insert_blank", after=0)
        self.assertEqual(self.pdf(d).page_count, 3)
        self.op(d, op="reorder", order=[2, 0, 1])
        self.op(d, op="rotate", pages=[0], angle=180)
        self.op(d, op="crop_pages", page=1, rect=[50, 50, 545, 792], pages=[1, 2])
        doc = self.pdf(d)
        self.assertEqual((round(doc[1].rect.width), round(doc[1].rect.height)), (495, 742))
        self.op(d, op="uncrop_pages", pages=[1, 2])
        self.assertEqual(round(self.pdf(d)[1].rect.width), 595)
        self.op(d, op="delete", pages=[0])
        self.assertEqual(self.pdf(d).page_count, 2)
        self.op(d, ok=False, op="delete", pages=[0, 1])
        # Insert another PDF after page 1, and an image file.
        r = self.c.post(f"/api/doc/{d}/merge", data={"file": [(io.BytesIO(form_pdf()), "f.pdf"),
                                                              (io.BytesIO(png_bytes()), "p.png")],
                                                     "after": "0"},
                        content_type="multipart/form-data")
        self.assertEqual(r.status_code, 200, r.get_json())
        self.assertEqual(len(r.get_json()["pages"]), 4)
        # Undo / redo.
        self.assertEqual(len(self.c.post(f"/api/doc/{d}/undo").get_json()["pages"]), 2)
        self.assertEqual(len(self.c.post(f"/api/doc/{d}/redo").get_json()["pages"]), 4)

    def test_split_and_extract(self):
        d = self.open()["id"]
        r = self.c.post(f"/api/doc/{d}/split", json={"ranges": "1, 2-"})
        names = zipfile.ZipFile(io.BytesIO(r.data)).namelist()
        self.assertEqual(names, ["test_p1.pdf", "test_p2.pdf"])
        self.assertEqual(self.c.post(f"/api/doc/{d}/split", json={"ranges": "5-9"}).status_code, 400)
        r = self.c.post(f"/api/doc/{d}/extract", json={"pages": [1]})
        self.assertEqual(pymupdf.open(stream=r.data).page_count, 1)

    def test_forms(self):
        d = self.open(form_pdf(), "form.pdf")["id"]
        widgets = self.c.get(f"/api/doc/{d}/page/0/widgets").get_json()
        name = next(w for w in widgets if w["name"] == "full_name")
        self.op(d, op="fill_form", values={str(name["xref"]): "Jane Doe"})
        r = self.op(d, op="fill_dates", value="10/04/2026")
        self.assertIn("1 date field", r.get_json()["message"])
        values = {w.field_name: w.field_value for w in self.pdf(d)[0].widgets()}
        self.assertEqual(values, {"full_name": "Jane Doe", "date_signed": "10/04/2026"})
        for kind in ("text", "multiline", "checkbox", "dropdown", "list"):
            self.op(d, op="add_field", page=0, type=kind, rect=[50, 300, 200, 320], choices="A, B")
        self.op(d, ok=False, op="add_field", page=0, type="dropdown", rect=[50, 300, 200, 320], choices="")
        self.op(d, ok=False, op="add_field", page=0, type="text", name="full_name", rect=[50, 300, 200, 320])
        self.assertEqual(len(list(self.pdf(d)[0].widgets())), 7)
        self.op(d, op="flatten")
        self.assertEqual(len(list(self.pdf(d)[0].widgets())), 0)

    def test_security_export(self):
        d = self.open()["id"]
        r = self.c.post(f"/api/doc/{d}/export", json={
            "userPassword": "open-me", "permissions": {"print": False, "copy": False}, "cleanMetadata": True})
        self.assertEqual(r.status_code, 200)
        out = pymupdf.open(stream=r.data)
        self.assertTrue(out.needs_pass)
        self.assertTrue(out.authenticate("open-me"))
        self.assertFalse(out.permissions & pymupdf.PDF_PERM_PRINT)
        self.assertFalse(out.permissions & pymupdf.PDF_PERM_COPY)
        self.assertTrue(out.permissions & pymupdf.PDF_PERM_ANNOTATE)
        self.assertNotIn("Secret", json.dumps(out.metadata))
        # Permissions without a password: opens freely but is restricted.
        r = self.c.post(f"/api/doc/{d}/export", json={"permissions": {"edit": False}})
        out = pymupdf.open(stream=r.data)
        self.assertFalse(out.needs_pass)
        self.assertIn("AES", out.metadata["encryption"])
        self.assertFalse(out.permissions & pymupdf.PDF_PERM_MODIFY)
        # Re-opening a protected file asks for the password.
        r = self.c.post("/api/open", data={"file": (io.BytesIO(
            self.c.post(f"/api/doc/{d}/export", json={"userPassword": "pw"}).data), "p.pdf")},
            content_type="multipart/form-data")
        self.assertEqual(r.status_code, 401)

    def test_metadata_watermark_compress(self):
        d = self.open()["id"]
        self.op(d, op="clean_metadata")
        meta = self.pdf(d).metadata
        self.assertFalse(meta.get("author") or meta.get("producer", "").startswith("Secret"))
        self.op(d, op="watermark", text="CONFIDENTIAL", angle=45, opacity=0.2)
        self.assertIn("CONFIDENTIAL", self.pdf(d)[1].get_text())
        self.op(d, op="watermark", dataUrl=data_url(png_bytes()), angle=0, opacity=0.3, pages=[0])
        r = self.op(d, op="watermark", text="سري", angle=45)
        self.assertIn("سري", text_of(self.pdf(d)[0]))
        self.op(d, ok=False, op="watermark", text="  ")
        r = self.op(d, op="compress", level="high")
        self.assertIn("Compressed", r.get_json()["message"])

    def test_text_layer(self):
        d = self.open()["id"]
        words = [{"text": "Scanned", "bbox": [100, 300, 180, 320]},
                 {"text": "words", "bbox": [185, 300, 230, 320]}]
        self.op(d, op="text_layer", pages=[{"page": 0, "words": words}, {"page": 1, "words": words}])
        for n in (0, 1):
            found = self.pdf(d)[n].search_for("Scanned words")
            self.assertTrue(found, f"page {n}")
            r = server.rect_to_display(self.pdf(d)[n], found[0])
            self.assertLess(abs(r[0] - 100), 3)
            self.assertLess(abs(r[1] - 300), 5)
        self.op(d, ok=False, op="text_layer", pages=[{"page": 0, "words": []}])

    def test_conversions_out(self):
        d = self.open()["id"]
        r = self.c.post(f"/api/doc/{d}/convert/docx")
        self.assertEqual(r.status_code, 200)
        import docx
        self.assertIn("Invoice", "\n".join(p.text for p in docx.Document(io.BytesIO(r.data)).paragraphs))
        r = self.c.post(f"/api/doc/{d}/convert/xlsx")
        from openpyxl import load_workbook
        self.assertTrue(load_workbook(io.BytesIO(r.data)).sheetnames)
        r = self.c.post(f"/api/doc/{d}/convert/pptx")
        from pptx import Presentation
        self.assertEqual(len(Presentation(io.BytesIO(r.data)).slides), 2)
        for fmt in ("png", "jpg", "webp"):
            r = self.c.post(f"/api/doc/{d}/convert/images", json={"format": fmt, "dpi": 72, "pages": [0]})
            self.assertEqual(r.status_code, 200)
            Image.open(io.BytesIO(r.data)).verify()
        r = self.c.post(f"/api/doc/{d}/convert/images", json={"format": "png", "dpi": 72})
        self.assertEqual(len(zipfile.ZipFile(io.BytesIO(r.data)).namelist()), 2)

    def test_conversions_in(self):
        import docx
        from openpyxl import Workbook
        from pptx import Presentation
        from pptx.util import Inches
        d = docx.Document()
        d.add_heading("Report title", 1)
        p = d.add_paragraph("Hello ")
        p.add_run("bold").bold = True
        d.add_paragraph("نص عربي")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text = "Cell A"
        b = io.BytesIO(); d.save(b)
        wb = Workbook(); wb.active.append(["Name", "Qty"]); wb.active.append(["Apples", 3])
        x = io.BytesIO(); wb.save(x)
        prs = Presentation(); s = prs.slides.add_slide(prs.slide_layouts[5]); s.shapes.title.text = "Slide one"
        s.shapes.add_picture(io.BytesIO(png_bytes()), Inches(1), Inches(2))
        pp = io.BytesIO(); prs.save(pp)
        cases = [("a.docx", b.getvalue(), ["Report title", "bold", "Cell A", "عربي"]),
                 ("b.xlsx", x.getvalue(), ["Apples"]),
                 ("c.pptx", pp.getvalue(), ["Slide one"]),
                 ("d.csv", b"x,y\n1,2\n", ["x"]),
                 ("e.txt", "plain text file".encode(), ["plain text file"])]
        for name, data, expect in cases:
            r = self.c.post("/api/import", data={"file": (io.BytesIO(data), name)},
                            content_type="multipart/form-data")
            self.assertEqual(r.status_code, 200, (name, r.get_json()))
            text = text_of(self.pdf(r.get_json()["id"])[0])
            for e in expect:
                self.assertIn(e, text, name)
        r = self.c.post("/api/import", data={"file": [(io.BytesIO(png_bytes()), "1.png"),
                                                       (io.BytesIO(png_bytes((60, 30))), "2.png")]},
                        content_type="multipart/form-data")
        self.assertEqual(len(r.get_json()["pages"]), 2)
        r = self.c.post("/api/import", data={"file": (io.BytesIO(b"x"), "x.exe")},
                        content_type="multipart/form-data")
        self.assertEqual(r.status_code, 400)

    def test_ai_without_key(self):
        d = self.open()["id"]
        import ai
        orig_key, orig_cfg = ai.saved_key, ai.CONFIG
        ai.CONFIG = Path(self.tmp.name) / "config.json"
        import os
        env = os.environ.pop("ANTHROPIC_API_KEY", None)
        try:
            self.assertFalse(self.c.get("/api/capabilities").get_json()["ai"]["configured"])
            r = self.c.post("/api/settings", json={"anthropicKey": "sk-ant-test"})
            self.assertTrue(r.get_json()["ai"]["configured"])
            self.assertEqual(ai.saved_key(), "sk-ant-test")
            self.c.post("/api/settings", json={"anthropicKey": ""})
            self.assertEqual(ai.saved_key(), "")
            r = self.c.post(f"/api/doc/{d}/ai/autofill", json={"info": "x"})
            self.assertEqual(r.status_code, 400)  # no form fields
            self.assertIn("no fillable", r.get_json()["error"])
            r = self.c.post(f"/api/doc/{d}/ai/chat", json={"messages": []})
            self.assertEqual(r.status_code, 400)
        finally:
            ai.CONFIG = orig_cfg
            if env is not None:
                os.environ["ANTHROPIC_API_KEY"] = env

    def test_version_changes_after_history_is_full(self):
        # Page images are cached by version, so it must change on every edit,
        # even once old undo steps are being dropped.
        d = self.open()["id"]
        seen = set()
        for i in range(server.MAX_HISTORY + 5):
            v = self.op(d, op="rect", page=0, rect=[10, 10, 20 + i, 20]).get_json()["version"]
            self.assertNotIn(v, seen)
            seen.add(v)
        v = self.c.post(f"/api/doc/{d}/undo").get_json()["version"]
        self.assertNotIn(v, seen)

    def test_unknown_op_and_bad_input(self):
        d = self.open()["id"]
        self.assertEqual(self.op(d, ok=False, op="nope").status_code, 400)
        self.assertEqual(self.op(d, ok=False, op="rotate").status_code, 400)  # missing keys
        self.assertEqual(self.c.get("/api/doc/missing").status_code, 404)



class HostedTest(unittest.TestCase):
    """PDFEDITOR_HOSTED=1: visitors are kept apart and server settings are locked."""

    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.saved = (server.WORKSPACE, server.HOSTED, server.MAX_DOCS_PER_OWNER, server.SHARE_AI_KEY)
        server.WORKSPACE = Path(self.tmp.name)
        server.HOSTED, server.MAX_DOCS_PER_OWNER, server.SHARE_AI_KEY = True, 2, False
        server.DOCS.clear()
        self.alice = server.app.test_client()
        self.bob = server.app.test_client()

    def tearDown(self):
        server.WORKSPACE, server.HOSTED, server.MAX_DOCS_PER_OWNER, server.SHARE_AI_KEY = self.saved
        server.DOCS.clear()
        self.tmp.cleanup()

    def open(self, client, data=None):
        return client.post("/api/open", data={"file": (io.BytesIO(data or sample_pdf()), "a.pdf")},
                           content_type="multipart/form-data")

    def test_visitors_are_isolated(self):
        r = self.open(self.alice)
        self.assertEqual(r.status_code, 200)
        doc_id = r.get_json()["id"]
        self.assertIn("pdfeditor_owner", r.headers.get("Set-Cookie", ""))
        self.assertEqual(len(self.alice.get("/api/docs").get_json()), 1)
        self.assertEqual(self.bob.get("/api/docs").get_json(), [])
        for method, url in (("get", f"/api/doc/{doc_id}"), ("get", f"/api/doc/{doc_id}/download"),
                            ("get", f"/api/doc/{doc_id}/page/0.png"), ("delete", f"/api/doc/{doc_id}")):
            self.assertEqual(getattr(self.bob, method)(url).status_code, 404, url)
        r = self.bob.post(f"/api/doc/{doc_id}/op", json={"op": "delete", "pages": [0]})
        self.assertEqual(r.status_code, 404)
        self.assertEqual(self.alice.get(f"/api/doc/{doc_id}").status_code, 200)

    def test_limits_and_locked_settings(self):
        self.assertEqual(self.open(self.alice).status_code, 200)
        self.assertEqual(self.open(self.alice).status_code, 200)
        r = self.open(self.alice)
        self.assertEqual(r.status_code, 400)
        self.assertIn("Close a tab", r.get_json()["error"])
        self.assertEqual(self.alice.post("/api/settings", json={"anthropicKey": "x"}).status_code, 403)
        r = self.alice.post("/api/fonts", data={"file": (io.BytesIO(b"x"), "f.ttf")},
                            content_type="multipart/form-data")
        self.assertEqual(r.status_code, 403)
        caps = self.alice.get("/api/capabilities").get_json()
        self.assertTrue(caps["hosted"])
        self.assertFalse(caps["ai"]["configured"])

    def test_ai_needs_visitor_key(self):
        doc_id = self.open(self.alice).get_json()["id"]
        r = self.alice.post(f"/api/doc/{doc_id}/ai/summary", json={})
        self.assertEqual(r.status_code, 400)
        self.assertIn("your browser", r.get_json()["error"])

    def test_idle_documents_are_forgotten(self):
        doc_id = self.open(self.alice).get_json()["id"]
        server.DOCS[doc_id]["touched"] -= 7 * 3600
        server._last_sweep[0] = 0
        saved = server.IDLE_HOURS
        server.IDLE_HOURS = 6
        try:
            self.assertEqual(self.alice.get("/api/docs").get_json(), [])
        finally:
            server.IDLE_HOURS = saved


if __name__ == "__main__":
    unittest.main()
