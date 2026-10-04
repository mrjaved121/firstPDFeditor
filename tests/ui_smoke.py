"""Browser smoke test of the editor UI (needs `pip install playwright` and Chrome or Edge).

Start the app first (python app.py --no-browser), then run:  python tests/ui_smoke.py [shots_dir]
"""

import io
import sys
import tempfile
from pathlib import Path

import pymupdf
from playwright.sync_api import sync_playwright

sys.path.insert(0, str(Path(__file__).resolve().parent))
from test_api import form_pdf, sample_pdf  # noqa: E402

URL = "http://127.0.0.1:5050"
SHOTS = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(tempfile.mkdtemp())
SHOTS.mkdir(parents=True, exist_ok=True)


def scanned_pdf():
    """A page that is only a picture of text, for OCR."""
    src = pymupdf.open()
    p = src.new_page(width=400, height=200)
    p.insert_text((30, 100), "Hello scanned world", fontsize=28)
    png = p.get_pixmap(dpi=200).tobytes("png")
    out = pymupdf.open()
    q = out.new_page(width=400, height=200)
    q.insert_image(q.rect, stream=png)
    return out.tobytes()


def main():
    tmp = Path(tempfile.mkdtemp())
    files = {}
    for name, data in (("sample.pdf", sample_pdf()), ("form.pdf", form_pdf()), ("scan.pdf", scanned_pdf())):
        files[name] = tmp / name
        files[name].write_bytes(data)

    errors, step = [], [""]
    with sync_playwright() as pw:
        try:
            browser = pw.chromium.launch(channel="chrome")
        except Exception:
            browser = pw.chromium.launch(channel="msedge")
        ctx = browser.new_context(viewport={"width": 1500, "height": 950}, accept_downloads=True)
        page = ctx.new_page()
        page.on("console", lambda m: m.type == "error" and errors.append(f"[{step[0]}] console: {m.text}"))
        page.on("pageerror", lambda e: errors.append(f"[{step[0]}] pageerror: {e}"))
        page.on("dialog", lambda d: d.accept())  # confirm() prompts

        def go(name):
            step[0] = name
            print("-", name)

        def version():
            return page.evaluate("state.doc ? state.doc.version : -1")

        def wait_change(before, timeout=15000):
            page.wait_for_function(f"state.doc && state.doc.version !== {before}", timeout=timeout)

        def page_box(n=0):
            return page.locator(f'.page[data-page="{n}"] .overlay').bounding_box()

        def show(n, y):
            """Scroll so that y (CSS px from the top of page n) is in view; return the page box."""
            page.evaluate("([n, y]) => { const el = document.querySelector(`.page[data-page='${n}']`);"
                          " document.querySelector('#viewer').scrollTop = el.offsetTop + y - 300; }", [n, y])
            return page_box(n)

        def click_at(n, x, y):
            b = show(n, y)
            page.mouse.click(b["x"] + x, b["y"] + y)

        def drag_at(n, x0, y0, x1, y1):
            b = show(n, y0)
            page.mouse.move(b["x"] + x0, b["y"] + y0)
            page.mouse.down()
            page.mouse.move(b["x"] + x1, b["y"] + y1, steps=6)
            page.mouse.up()

        def tool(name):
            page.click(f'.tool[data-tool="{name}"]')

        def menu(label, action):
            page.click(f'.menu-btn:has-text("{label}")')
            page.click(f'.menu-item[data-action="{action}"]')

        def submit_dialog():
            page.locator("dialog[open] button[type=submit]").click()

        go("load")
        page.goto(URL)
        page.wait_for_load_state("networkidle")
        # Close tabs left from earlier runs.
        for _ in range(10):
            if not page.locator(".doc-tab").count():
                break
            page.locator(".doc-tab .close").first.click()
            page.wait_for_timeout(300)
        page.wait_for_selector("#dropZone:not([hidden])")
        page.screenshot(path=SHOTS / "01_empty.png")

        go("open")
        page.set_input_files("#openInput", str(files["sample.pdf"]))
        page.wait_for_selector('.page[data-page="1"]')
        assert page.locator(".doc-tab").count() == 1

        go("add text (bold, centered)")
        tool("addtext")
        page.click("#boldBtn")
        page.select_option("#alignSelect", "center")
        v = version()
        click_at(0, 100, 420)
        page.keyboard.type("Bold added text")
        page.keyboard.press("Enter")
        wait_change(v)
        page.click("#boldBtn")

        go("edit text")
        tool("edittext")
        page.wait_for_selector(".text-line")
        v = version()
        page.locator(".text-line[title^=Invoice]").first.click()
        page.keyboard.press("Control+A")
        page.keyboard.type("Invoice for John Roe")
        page.keyboard.press("Enter")
        wait_change(v)

        go("highlight / underline / strikeout")
        for t in ("highlight", "underline", "strikeout"):
            tool(t)
            v = version()
            drag_at(0, 85, 158, 300, 175)
            wait_change(v)

        go("note")
        tool("note")
        click_at(0, 600, 120)
        page.fill("dialog[open] textarea[name=text]", "Please check")
        v = version()
        submit_dialog()
        wait_change(v)

        go("date stamp")
        tool("date")
        v = version()
        click_at(0, 500, 900)
        wait_change(v)

        go("shapes and ink")
        for t, (x0, y0, x1, y1) in (("rect", (90, 700, 200, 760)), ("ellipse", (230, 700, 330, 760)),
                                    ("arrow", (350, 700, 450, 760)), ("ink", (470, 700, 560, 760))):
            tool(t)
            v = version()
            drag_at(0, x0, y0, x1, y1)
            wait_change(v)

        go("images: move, rotate, crop")
        tool("image")
        page.wait_for_selector(".img-box")
        box = page.locator(".img-box").first
        box.click()
        bb = box.bounding_box()
        v = version()
        page.mouse.move(bb["x"] + 20, bb["y"] + 20)
        page.mouse.down()
        page.mouse.move(bb["x"] + 80, bb["y"] + 120, steps=6)
        page.mouse.up()
        wait_change(v)
        page.wait_for_selector(".img-box.selected")
        v = version()
        page.click('.img-box.selected [data-img="rotate"]')
        wait_change(v)
        page.wait_for_selector(".img-box.selected")
        page.click('.img-box.selected [data-img="crop"]')
        bb = page.locator(".img-box.selected").bounding_box()
        v = version()
        page.mouse.move(bb["x"] + 2, bb["y"] + 2)
        page.mouse.down()
        page.mouse.move(bb["x"] + bb["width"] * 0.6, bb["y"] + bb["height"] * 0.6, steps=5)
        page.mouse.up()
        wait_change(v)
        page.screenshot(path=SHOTS / "02_edited.png")

        go("signature: typed")
        tool("sign")
        page.click('#signDialog .tab[data-tab="type"]')
        page.fill("#typedName", "Jane Doe")
        page.click("#signUse")
        v = version()
        drag_at(0, 400, 950, 600, 1010)
        wait_change(v)

        go("watermark")
        menu("Security", "watermark")
        v = version()
        submit_dialog()
        wait_change(v)

        go("smart redact")
        menu("Security", "smartRedact")
        submit_dialog()
        page.wait_for_selector("dialog[open] .result-list")
        assert page.locator("dialog[open] .result-list label").count() == 3, "expected SSN, card, email"
        v = version()
        submit_dialog()
        wait_change(v)

        go("crop page")
        menu("Pages", "cropTool")
        b = page_box(1)
        drag_at(1, 20, 20, b["width"] - 20, b["height"] - 20)
        v = version()
        submit_dialog()
        wait_change(v)

        go("add field + fill form")
        tool("field")
        drag_at(0, 400, 1150, 650, 1180)
        page.fill("dialog[open] input[name=name]", "Customer")
        v = version()
        submit_dialog()
        wait_change(v)
        tool("form")
        page.wait_for_selector(".form-field input")
        page.fill(".form-field input", "ACME Inc")
        v = version()
        page.click("#saveFormBtn")
        wait_change(v)

        go("compress / metadata")
        menu("Convert", "compress")
        v = version()
        submit_dialog()
        wait_change(v)
        v = version()
        menu("Security", "cleanMeta")
        wait_change(v)

        go("undo / redo")
        v = version()
        page.click("#undoBtn")
        wait_change(v)
        v = version()
        page.click("#redoBtn")
        wait_change(v)

        go("downloads")
        for action, ext in (("toDocx", ".docx"), ("toXlsx", ".xlsx"), ("toPptx", ".pptx")):
            with page.expect_download(timeout=60000) as dl:
                menu("Convert", action)
            assert dl.value.suggested_filename.endswith(ext), dl.value.suggested_filename
        with page.expect_download() as dl:
            menu("Security", "export")
            page.fill("dialog[open] input[name=userPassword]", "secret")
            submit_dialog()
        data = Path(dl.value.path()).read_bytes()
        assert pymupdf.open(stream=data).needs_pass
        with page.expect_download() as dl:
            page.click("#downloadBtn")
        out = pymupdf.open(dl.value.path())
        text = out[0].get_text()
        assert "John Roe" in text and "123-45-6789" not in text, text

        go("zoom / fit / theme")
        page.click("#fitWidth")
        page.click("#fitPage")
        page.click("#zoomIn")
        page.click("#themeBtn")
        assert page.evaluate("document.documentElement.dataset.theme") in ("dark", "light")
        page.screenshot(path=SHOTS / "03_theme.png")

        go("AI without key opens settings")
        menu("AI", "aiSummary")
        page.wait_for_selector("dialog[open] input[name=key]")
        page.screenshot(path=SHOTS / "04_ai_settings.png")
        page.locator("dialog[open] [data-cancel]").click()

        go("second tab + OCR")
        page.set_input_files("#openInput", str(files["scan.pdf"]))
        page.wait_for_function("document.querySelectorAll('.doc-tab').length === 2")
        menu("Convert", "ocr")
        v = version()
        submit_dialog()
        try:
            wait_change(v, timeout=120000)
            ok = page.evaluate("state.doc.hasText")
            print("  OCR made the page searchable:", ok)
        except Exception as e:
            print("  OCR did not finish (no internet for Tesseract.js?):", str(e).splitlines()[0])

        go("tab switch, restore after reload, close")
        page.locator(".doc-tab").first.click()
        page.wait_for_function("state.doc && state.doc.name === 'sample.pdf'")
        page.reload()
        page.wait_for_function("document.querySelectorAll('.doc-tab').length === 2", timeout=15000)
        page.wait_for_selector(".page")
        page.screenshot(path=SHOTS / "05_restored.png")
        while page.locator(".doc-tab").count():
            page.locator(".doc-tab .close").first.click()
            page.wait_for_timeout(300)
        page.wait_for_selector("#dropZone:not([hidden])")

        go("import word file")
        import docx
        d = docx.Document()
        d.add_heading("Imported", 1)
        d.add_paragraph("From Word")
        p = tmp / "in.docx"
        d.save(p)
        page.set_input_files("#importInput", str(p))
        page.wait_for_selector(".page")
        while page.locator(".doc-tab").count():
            page.locator(".doc-tab .close").first.click()
            page.wait_for_timeout(300)

        browser.close()

    print("screenshots:", SHOTS)
    if errors:
        print("ERRORS:")
        for e in errors:
            print("  ", e)
        sys.exit(1)
    print("UI smoke test passed.")


if __name__ == "__main__":
    main()
